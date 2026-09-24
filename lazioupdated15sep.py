#!/usr/bin/env python3
"""
Fast Regione Lazio - Bandi di gara in scadenza scraper.

Scrapes all tender list records and extracts 'Data Pubblicazione'
from the 'TABELLA INFORMATIVA DI INDICIZZAZIONE' section on detail pages.
"""

from __future__ import annotations

import argparse
import asyncio
import html
import re
import time
from collections import OrderedDict
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from playwright.async_api import Page, async_playwright, TimeoutError as PlaywrightTimeoutError

# -----------------------------------------------------------------------------
# Configuration
# -----------------------------------------------------------------------------
BASE_URL = "https://centraleacquisti.regione.lazio.it"
LIST_URL = BASE_URL + "/bandi-e-strumenti-di-acquisto/bandi-di-gara-in-scadenza"

OUTPUT_DIR = Path(__file__).with_name("output")
CONTRACT_XLSX = OUTPUT_DIR / "lazio_contracts.xlsx"

REQUIRED_CONTRACT_COLUMNS = [
    "TIPO",
    "CIG/N.Gara",
    "TITOLO",
    "ENTE",
    "IMPORTO",
    "SCADENZA",
    "Data Pubblicazione",
]

CONTRACT_EXTRA_COLUMNS = [
    "id_doc",
    "tipo_doc",
    "URL dettaglio",
]

REQUEST_TIMEOUT_MS = 60_000
NAV_TIMEOUT_MS = 90_000
PAGE_PAUSE_S = 0.5

CONTRACT_HEADER_ALIASES = {
    "TIPO": ["tipo", "tipologia", "tipo procedura"],
    "CIG/N.Gara": ["cig/n.gara", "cig / n.gara", "cig/n. gara", "cig", "n.gara", "n. gara", "numero gara",
                   "registro di sistema"],
    "TITOLO": ["titolo", "oggetto", "descrizione"],
    "ENTE": ["ente", "amministrazione", "stazione appaltante"],
    "IMPORTO": ["importo", "valore"],
    "SCADENZA": ["scadenza", "termine presentazione offerta", "termine offerte"],
    "Data Pubblicazione": ["data pubblicazione", "data di pubblicazione", "pubblicazione"],
}


# -----------------------------------------------------------------------------
# Text / Name Helpers
# -----------------------------------------------------------------------------
def clean_text(value) -> str:
    if value is None:
        return ""
    text = html.unescape(str(value)).replace("\xa0", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return text.strip()


def simplify(value: str) -> str:
    value = clean_text(value).lower()
    value = re.sub(r"[^a-z0-9àèéìòù/]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def extract_id_doc(url: str) -> str:
    try:
        return parse_qs(urlsplit(url).query).get("id_doc", [""])[0]
    except Exception:
        return ""


def extract_tipo_doc(url: str) -> str:
    try:
        return parse_qs(urlsplit(url).query).get("tipo_doc", [""])[0]
    except Exception:
        return ""


# -----------------------------------------------------------------------------
# Browser Helpers & List Extraction
# -----------------------------------------------------------------------------
async def wait_for_render(page: Page) -> None:
    try:
        await page.wait_for_load_state("domcontentloaded", timeout=NAV_TIMEOUT_MS)
    except PlaywrightTimeoutError:
        pass
    try:
        await page.wait_for_load_state("networkidle", timeout=10_000)
    except PlaywrightTimeoutError:
        pass
    await page.wait_for_timeout(500)


async def dismiss_cookie_banner(page: Page) -> None:
    candidates = [
        "button:has-text('Accetta')",
        "button:has-text('Accetto')",
        "button:has-text('Accept')",
        "button:has-text('Consenti')",
        "#onetrust-accept-btn-handler",
    ]
    for sel in candidates:
        try:
            loc = page.locator(sel).first
            if await loc.count() and await loc.is_visible(timeout=500):
                await loc.click(timeout=1500)
                await page.wait_for_timeout(300)
                return
        except Exception:
            continue


def _map_contract_headers(headers: list[str]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    simple_headers = [simplify(h) for h in headers]
    for output_name, aliases in CONTRACT_HEADER_ALIASES.items():
        for idx, h in enumerate(simple_headers):
            if any(simplify(alias) == h or simplify(alias) in h for alias in aliases):
                mapped[output_name] = idx
                break
    return mapped


# The current Angular table uses app-tabella-td wrappers around td cells,
# a header row of td elements, and buttons (not anchors) for detail navigation.
TABLE = 'table:has-text("CIG/N.Gara"):has-text("SCADENZA")'
RESULT_TIMEOUT_S = 180


async def table_rows(page):
    table = page.locator(TABLE + ":visible").first
    if not await table.count():
        return []
    return await table.locator("tr").evaluate_all("""
        rows => rows.map(row => Array.from(row.querySelectorAll('td, th'))
          .filter(cell => cell.closest('tr') === row)
          .map(cell => (cell.innerText || '').trim()))
          .filter(cells => cells.length >= 6 && cells[0] !== 'TIPO')
    """)


async def wait_for_results(page, previous=None):
    deadline = time.monotonic() + RESULT_TIMEOUT_S
    last = None
    stable_since = time.monotonic()
    while time.monotonic() < deadline:
        rows = await table_rows(page)
        loading = await page.get_by_text('Caricamento...', exact=True).is_visible()
        if rows != last:
            last, stable_since = rows, time.monotonic()
        if rows and rows != previous and not loading and time.monotonic() - stable_since >= 1:
            return rows
        await asyncio.sleep(0.5)
    raise RuntimeError('Tender rows did not load/change within 180 seconds. See output diagnostics.')


async def capture_detail_url(page, row_index):
    # Observe the real navigation URL generated by the button, while keeping
    # the list loaded on its current page. Detail pages are visited afterwards.
    found = asyncio.get_running_loop().create_future()
    context = page.context
    existing = set(context.pages)

    active_handlers = set()

    async def intercept(route):
        task = asyncio.current_task()
        active_handlers.add(task)
        try:
            request = route.request
            if request.is_navigation_request() and extract_id_doc(request.url):
                # Finish handling the request BEFORE waking the caller.
                # Otherwise its cleanup can unroute/close a popup mid-abort.
                await route.abort('aborted')
                if not found.done():
                    found.set_result(request.url)
            else:
                await route.continue_()
        except Exception as exc:
            # Propagate through the awaited future, not an event-listener error.
            if not found.done():
                found.set_exception(exc)
        finally:
            active_handlers.discard(task)

    pattern = '**/*id_doc=*'
    await context.route(pattern, intercept)
    try:
        table = page.locator(TABLE + ":visible").first
        row = table.locator('tr').filter(has=page.locator('button')).nth(row_index)
        try:
            await row.locator('button').first.click(timeout=15_000, no_wait_after=True)
        except PlaywrightTimeoutError:
            if not found.done():
                raise
        return await asyncio.wait_for(found, timeout=30)
    finally:
        await context.unroute(pattern, intercept)
        # unroute does not wait for callbacks that are already running.
        if active_handlers:
            await asyncio.gather(*tuple(active_handlers))
        for opened in list(context.pages):
            if opened not in existing:
                await opened.close()


async def save_diagnostics(page):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    try:
        (OUTPUT_DIR / 'lazio_debug.html').write_text(await page.content(), encoding='utf-8')
        await page.screenshot(path=str(OUTPUT_DIR / 'lazio_debug.png'), full_page=True)
    except Exception as exc:
        print(f'Diagnostics could not be saved: {exc}')


async def scrape_all_contracts(page: Page, max_pages: int | None = None):
    print(f'[1] Opening {LIST_URL}', flush=True)
    await page.goto(LIST_URL, wait_until='domcontentloaded', timeout=NAV_TIMEOUT_MS)
    await dismiss_cookie_banner(page)
    rows_out = OrderedDict()
    seen_pages = set()
    page_no = 1
    try:
        rows = await wait_for_results(page)
        while True:
            signature = tuple(tuple(r) for r in rows)
            if signature in seen_pages:
                raise RuntimeError(f'Pagination repeated a previous page at page {page_no}.')
            seen_pages.add(signature)
            table = page.locator(TABLE + ":visible").first
            header = await table.locator('tr').first.locator('td, th').all_inner_texts()
            mapping = _map_contract_headers(header)
            if not all(c in mapping for c in REQUIRED_CONTRACT_COLUMNS[:6]):
                raise RuntimeError(f'Unexpected table headers: {header}')
            buttons = table.locator('tr').filter(has=page.locator('button'))
            if await buttons.count() != len(rows):
                raise RuntimeError('Row/button counts differ; stopping to avoid assigning wrong URLs.')
            print(f'    Page {page_no}: {len(rows)} results', flush=True)
            for idx, cells in enumerate(rows):
                url = await capture_detail_url(page, idx)
                rec = {c: '' for c in REQUIRED_CONTRACT_COLUMNS}
                for col, position in mapping.items():
                    rec[col] = clean_text(cells[position]) if position < len(cells) else ''
                rec.update({'URL dettaglio': url, 'id_doc': extract_id_doc(url),
                            'tipo_doc': extract_tipo_doc(url)})
                key = (rec['tipo_doc'], rec['id_doc'])
                rows_out[key] = rec
                print(f"      {idx+1}/{len(rows)}: id_doc={rec['id_doc']}", flush=True)
            # Checkpoint every page, before detail enrichment.
            write_excel(CONTRACT_XLSX, list(rows_out.values()), REQUIRED_CONTRACT_COLUMNS + CONTRACT_EXTRA_COLUMNS)
            if max_pages and page_no >= max_pages:
                print('Stopped at requested --max-pages limit.')
                break
            next_button = page.locator('button:visible').filter(has_text=re.compile(r'^\s*' + str(page_no+1) + r'\s*$')).first
            if not await next_button.count() or not await next_button.is_enabled():
                totals = await page.get_by_text(re.compile(r'Risultati totali')).all_inner_texts()
                total_match = re.search(r'Risultati totali\s*([\d.,]+)', ' '.join(totals))
                expected = int(re.sub(r'\D', '', total_match[1])) if total_match else None
                if expected is None or sum(len(p) for p in seen_pages) < expected:
                    raise RuntimeError(f'Next page button not found; collected {len(rows_out)}, portal total={expected}.')
                print(f'    Finished: {len(rows_out)} unique tenders.', flush=True)
                break
            await next_button.click(timeout=15_000)
            rows = await wait_for_results(page, previous=rows)
            page_no += 1
    except Exception:
        if rows_out:
            partial_path = OUTPUT_DIR / 'lazio_contracts_partial.xlsx'
            write_excel(partial_path, list(rows_out.values()), REQUIRED_CONTRACT_COLUMNS + CONTRACT_EXTRA_COLUMNS)
            print(f'Incomplete run; partial results saved to {partial_path}')
        await save_diagnostics(page)
        raise
    return list(rows_out.values())


# -----------------------------------------------------------------------------
# Detailed Page Scraper for Data Pubblicazione
# -----------------------------------------------------------------------------
async def extract_publication_date_from_detail(page: Page, url: str) -> str:
    try:
        await page.goto(url, wait_until="domcontentloaded", timeout=NAV_TIMEOUT_MS)
        await wait_for_render(page)

        accordion_button = page.get_by_role('button', name=re.compile('TABELLA INFORMATIVA DI INDICIZZAZIONE', re.I)).first
        try:
            await accordion_button.wait_for(state='visible', timeout=60_000)
            if await accordion_button.get_attribute('aria-expanded') != 'true':
                await accordion_button.click()
            await page.get_by_text(re.compile(r'Data\s+(?:di\s+)?Pubblicazione', re.I)).first.wait_for(timeout=30_000)
        except PlaywrightTimeoutError:
            pass

        # Parse text from rendered DOM
        pub_date = await page.evaluate(
            """
            () => {
                const clean = s => (s || '').replace(/\\s+/g, ' ').trim();

                // 1. Look specifically within tables or containers near 'TABELLA INFORMATIVA'
                const headings = Array.from(document.querySelectorAll('*'));
                const section = headings.find(h => clean(h.innerText).toUpperCase().includes('TABELLA INFORMATIVA DI INDICIZZAZIONE'));

                let searchText = document.body.innerText;
                if (section) {
                    const container = section.closest('div, section, fieldset') || section.parentElement;
                    if (container) searchText = container.innerText;
                }

                // 2. Extract date pattern after "Data Pubblicazione"
                const match = searchText.match(/Data\\s+(?:di\\s+)?Pubblicazione\\s*:?\\s*([0-3]?\\d\\/[01]?\\d\\/(?:19|20)\\d{2})/i);
                if (match) return match[1];

                // 3. Fallback: Check table rows directly
                const rows = Array.from(document.querySelectorAll('tr, .row, .form-group'));
                for (const row of rows) {
                    const txt = clean(row.innerText);
                    if (txt.toLowerCase().includes('data pubblicazione') || txt.toLowerCase().includes('data di pubblicazione')) {
                        const dateMatch = txt.match(/([0-3]?\\d\\/[01]?\\d\\/(?:19|20)\\d{2})/);
                        if (dateMatch) return dateMatch[1];
                    }
                }
                return "";
            }
            """
        )
        return pub_date
    except Exception as exc:
        print(f"      Failed to load detail page ({url}): {exc}")
        return ""


async def populate_publication_dates(page: Page, contracts: list[dict[str, str]]):
    print(f"[2] Opening detail pages to extract 'Data Pubblicazione' for {len(contracts)} contracts...")
    for idx, contract in enumerate(contracts, 1):
        url = contract["URL dettaglio"]
        doc_id = contract.get("id_doc", "")
        print(f"    [{idx}/{len(contracts)}] Processing tender id_doc={doc_id}...")

        date = await extract_publication_date_from_detail(page, url)
        contract["Data Pubblicazione"] = date
        print(f"      -> Data Pubblicazione: {date if date else 'NOT FOUND'}")


# -----------------------------------------------------------------------------
# Excel Output
# -----------------------------------------------------------------------------
def write_excel(path: Path, rows: list[dict[str, str]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Dati"
    ws.append(columns)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([clean_text(row.get(c, "")) for c in columns])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for idx, col in enumerate(columns, 1):
        max_len = len(col)
        for row in rows[:500]:
            first_line = clean_text(row.get(col, "")).split("\n", 1)[0]
            max_len = max(max_len, min(len(first_line), 60))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 10), 62)

    wb.save(path)
    print(f"[3] Excel saved: {path} ({len(rows)} contracts)")


# -----------------------------------------------------------------------------
# Main Execution
# -----------------------------------------------------------------------------
def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fast Scraper for Regione Lazio Tenders")
    parser.add_argument("--headed", action="store_true", help="Show Chromium browser UI")
    parser.add_argument("--max-pages", type=int, default=None, help="Limit scraped list pages")
    parser.add_argument('--list-only', action='store_true', help='Collect all list rows and detail URLs; skip publication-date visits')
    args = parser.parse_args()
    if args.max_pages is not None and args.max_pages < 1:
        parser.error('--max-pages must be at least 1')
    return args


async def main_async() -> None:
    args = parse_args()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=not args.headed)
        context = await browser.new_context(
            locale="it-IT",
        )
        page = await context.new_page()
        page.set_default_timeout(REQUEST_TIMEOUT_MS)

        try:
            contracts = await scrape_all_contracts(page, max_pages=args.max_pages)
            if contracts:
                if not args.list_only:
                    await populate_publication_dates(page, contracts)
                write_excel(CONTRACT_XLSX, contracts, REQUIRED_CONTRACT_COLUMNS + CONTRACT_EXTRA_COLUMNS)
            else:
                print("No contracts collected.")
        finally:
            await context.close()
            await browser.close()


def main():
    asyncio.run(main_async())


if __name__ == "__main__":
    main()